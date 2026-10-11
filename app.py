import streamlit as st
import os
import sys
import types
import pandas as pd

# 1. Page Configuration (Rendered FIRST so the screen is never blank!)
st.set_page_config(page_title="BraiNode: Pro Academic RAG", page_icon="🧠", layout="wide")

st.title("🧠 BraiNode: Hybrid Academic RAG & Study Suite")
st.markdown(
    "**Features:** Hybrid Search (BM25 + Vector) • Table-Aware Multi-Format Parsing (`.pdf`, `.docx`, `.txt`) "
    "• Automated Viva/Study Tools • RAGAS Quantitative Evaluation"
)

# --- HOTFIX FOR RAGAS BUG ---
dummy_chat = types.ModuleType("langchain_community.chat_models.vertexai")
dummy_chat.ChatVertexAI = type("ChatVertexAI", (object,), {})
sys.modules["langchain_community.chat_models.vertexai"] = dummy_chat
# ----------------------------

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_qdrant import QdrantVectorStore
from langchain_community.retrievers import BM25Retriever
from langchain_classic.retrievers import EnsembleRetriever
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_classic.chains import create_retrieval_chain, create_history_aware_retriever
from langchain_classic.chains.combine_documents import create_stuff_documents_chain
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_core.messages import HumanMessage, AIMessage
from langchain_core.documents import Document

import pdfplumber
import docx

# 2. Cached Embedding Model (Fixes slow startup & reloads)
@st.cache_resource(show_spinner="Loading local AI embedding model into memory...")
def load_embeddings():
    return HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")

# 3. Table-Aware Multi-Format Document Parser
def parse_uploaded_files(uploaded_files):
    docs = []
    for uploaded_file in uploaded_files:
        file_name = uploaded_file.name
        ext = os.path.splitext(file_name)[1].lower()
        
        if ext == ".pdf":
            with pdfplumber.open(uploaded_file) as pdf:
                for page_num, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    
                    # Extract tables and convert them into Markdown format for Gemini
                    tables = page.extract_tables()
                    table_md = ""
                    for table in tables:
                        cleaned_rows = [
                            [str(cell).replace("\n", " ").strip() if cell else "" for cell in row]
                            for row in table if any(row)
                        ]
                        if cleaned_rows:
                            table_md += "\n\n**[Extracted Table]:**\n"
                            for row in cleaned_rows:
                                table_md += "| " + " | ".join(row) + " |\n"
                    
                    full_content = f"{text}\n{table_md}".strip()
                    if full_content:
                        docs.append(Document(
                            page_content=full_content,
                            metadata={"source": file_name, "page": page_num}
                        ))
                        
        elif ext == ".docx":
            doc = docx.Document(uploaded_file)
            full_text = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells])
                    if row_text:
                        full_text.append(f"| {row_text} |")
            if full_text:
                docs.append(Document(
                    page_content="\n".join(full_text),
                    metadata={"source": file_name, "page": 1}
                ))
                
        elif ext == ".txt":
            text = uploaded_file.read().decode("utf-8", errors="ignore")
            if text.strip():
                docs.append(Document(
                    page_content=text,
                    metadata={"source": file_name, "page": 1}
                ))
    return docs

# 4. Sidebar Setup
st.sidebar.header("1. API Configuration")
api_key = st.sidebar.text_input("Enter Google Gemini API Key:", type="password")

st.sidebar.header("2. Document Management")
uploaded_files = st.sidebar.file_uploader(
    "Upload PDF, DOCX, or TXT files", 
    type=["pdf", "docx", "txt"], 
    accept_multiple_files=True
)

search_mode = st.sidebar.radio(
    "Retrieval Mode:",
    ["Hybrid Search (BM25 + Qdrant)", "Dense Vector Only (Qdrant)"],
    index=0
)

process_btn = st.sidebar.button("⚡ Process & Index Documents")
clear_btn = st.sidebar.button("🗑️ Clear Session")

# Initialize Session States
if "chat_history" not in st.session_state or clear_btn:
    st.session_state.chat_history = []
if "langchain_history" not in st.session_state or clear_btn:
    st.session_state.langchain_history = []
if "rag_chain" not in st.session_state or clear_btn:
    st.session_state.rag_chain = None
if "raw_chunks" not in st.session_state or clear_btn:
    st.session_state.raw_chunks = []
if "eval_data" not in st.session_state or clear_btn:
    st.session_state.eval_data = {"question": [], "answer": [], "contexts": []}
if "study_notes" not in st.session_state or clear_btn:
    st.session_state.study_notes = ""

# 5. Hybrid Vectorization Pipeline
if process_btn:
    if not api_key:
        st.sidebar.error("Please enter your Gemini API key first!")
    elif not uploaded_files:
        st.sidebar.error("Please upload at least one document!")
    else:
        os.environ["GOOGLE_API_KEY"] = api_key
        with st.spinner("Extracting text, parsing tables, and building Hybrid Index..."):
            documents = parse_uploaded_files(uploaded_files)
            
            if not documents:
                st.sidebar.error("Could not extract any readable text from the uploaded files.")
            else:
                text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)
                chunks = text_splitter.split_documents(documents)
                st.session_state.raw_chunks = chunks

                # Load cached embeddings
                embeddings = load_embeddings()
                
                # 1. Dense Vector Retriever (Qdrant)
                vector_store = QdrantVectorStore.from_documents(
                    chunks,
                    embedding=embeddings,
                    location=":memory:",
                    collection_name="eval_docs"
                )
                qdrant_retriever = vector_store.as_retriever(search_kwargs={"k": 3})

                # 2. Sparse Keyword Retriever (BM25) + Ensemble Hybrid
                if "Hybrid" in search_mode:
                    bm25_retriever = BM25Retriever.from_documents(chunks)
                    bm25_retriever.k = 3
                    retriever = EnsembleRetriever(
                        retrievers=[bm25_retriever, qdrant_retriever],
                        weights=[0.5, 0.5]
                    )
                else:
                    retriever = qdrant_retriever

                llm = ChatGoogleGenerativeAI(model="gemini-1.5-flash", temperature=0.3)

                contextualize_q_system_prompt = (
                    "Given a chat history and the latest user question "
                    "which might reference context in the chat history, "
                    "formulate a standalone question which can be understood "
                    "without the chat history. Do NOT answer the question, "
                    "just reformulate it if needed and otherwise return it as is."
                )
                contextualize_q_prompt = ChatPromptTemplate.from_messages([
                    ("system", contextualize_q_system_prompt),
                    MessagesPlaceholder("chat_history"),
                    ("human", "{input}"),
                ])
                history_aware_retriever = create_history_aware_retriever(
                    llm, retriever, contextualize_q_prompt
                )

                qa_system_prompt = (
                    "You are BraiNode, an advanced academic AI assistant. "
                    "Use the following pieces of retrieved context (including any extracted tables) "
                    "to answer the question accurately. Cite the source document or page number "
                    "when possible. If you don't know the answer, state that clearly.\n\n"
                    "{context}"
                )
                qa_prompt = ChatPromptTemplate.from_messages([
                    ("system", qa_system_prompt),
                    MessagesPlaceholder("chat_history"),
                    ("human", "{input}"),
                ])
                
                question_answer_chain = create_stuff_documents_chain(llm, qa_prompt)
                st.session_state.rag_chain = create_retrieval_chain(history_aware_retriever, question_answer_chain)
                
                st.sidebar.success(f"Indexed {len(chunks)} chunks across {len(uploaded_files)} file(s) using {search_mode}!")

# 6. Interactive 3-Tab Interface
if not api_key:
    st.info("👈 Enter your Google Gemini API Key in the sidebar and upload a document to begin.")
else:
    os.environ["GOOGLE_API_KEY"] = api_key

chat_tab, study_tab, eval_tab = st.tabs([
    "💬 Hybrid RAG Chat", 
    "🎓 Study Tools & Export", 
    "📊 RAGAS Evaluation"
])

# --- TAB 1: CHAT INTERFACE ---
with chat_tab:
    for message in st.session_state.chat_history:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    if user_query := st.chat_input("Ask about concepts, equations, or tables in your documents..."):
        if not api_key:
            st.error("⚠️ Please enter your Gemini API key in the sidebar first!")
        elif not st.session_state.rag_chain:
            st.error("⚠️ Please upload and process a document using the sidebar first!")
        else:
            st.session_state.chat_history.append({"role": "user", "content": user_query})
            with st.chat_message("user"):
                st.markdown(user_query)

            with st.chat_message("assistant"):
                with st.spinner("Running Hybrid Search (BM25 + Vector)..."):
                    response = st.session_state.rag_chain.invoke({
                        "input": user_query,
                        "chat_history": st.session_state.langchain_history
                    })
                    
                    answer = response["answer"]
                    st.markdown(answer)

                    # Show retrieved source chunks in an expander for transparency
                    with st.expander("🔍 View Retrieved Context & Sources"):
                        for idx, doc in enumerate(response["context"], 1):
                            src = doc.metadata.get("source", "Unknown")
                            pg = doc.metadata.get("page", "?")
                            st.markdown(f"**Chunk {idx} — {src} (Page {pg}):**")
                            st.caption(doc.page_content[:500] + "...")

                    retrieved_contexts = [doc.page_content for doc in response["context"]]
                    st.session_state.eval_data["question"].append(user_query)
                    st.session_state.eval_data["answer"].append(answer)
                    st.session_state.eval_data["contexts"].append(retrieved_contexts)

            st.session_state.chat_history.append({"role": "assistant", "content": answer})
            st.session_state.langchain_history.extend([
                HumanMessage(content=user_query),
                AIMessage(content=answer)
            ])

# --- TAB 2: AUTOMATED STUDY TOOLS & EXPORT ---
with study_tab:
    st.header("🎓 Automated Academic Study Suite")
    st.markdown("Generate instant revision notes, viva questions, or export your chat log for your project report.")
    
    col1, col2 = st.columns(2)
    
    with col1:
        if st.button("📝 Generate Executive Summary & Key Takeaways"):
            if not st.session_state.raw_chunks:
                st.warning("Please upload and process a document first!")
            else:
                with st.spinner("Synthesizing document summary..."):
                    llm = ChatGoogleGenerativeAI(model="gemini-1.5-flash", temperature=0.3)
                    sample_text = "\n\n".join([c.page_content for c in st.session_state.raw_chunks[:12]])
                    prompt = (
                        "Provide a structured academic Executive Summary of the following document text. "
                        "Include: 1. Core Objective, 2. Key Concepts/Definitions, and 3. Major Takeaways.\n\n"
                        f"{sample_text}"
                    )
                    res = llm.invoke(prompt)
                    st.session_state.study_notes = res.content

    with col2:
        if st.button("🎯 Generate 5 Viva / Exam Questions with Answers"):
            if not st.session_state.raw_chunks:
                st.warning("Please upload and process a document first!")
            else:
                with st.spinner("Creating viva questions and model answers..."):
                    llm = ChatGoogleGenerativeAI(model="gemini-1.5-flash", temperature=0.4)
                    sample_text = "\n\n".join([c.page_content for c in st.session_state.raw_chunks[:12]])
                    prompt = (
                        "Act as a university examiner. Based on the following document text, generate "
                        "5 challenging Viva / Exam questions along with detailed model answers.\n\n"
                        f"{sample_text}"
                    )
                    res = llm.invoke(prompt)
                    st.session_state.study_notes = res.content

    if st.session_state.study_notes:
        st.markdown("---")
        st.markdown(st.session_state.study_notes)
        st.download_button(
            label="📥 Download Study Notes (.md)",
            data=st.session_state.study_notes,
            file_name="BraiNode_Study_Notes.md",
            mime="text/markdown"
        )

    if st.session_state.chat_history:
        st.markdown("---")
        chat_export = "# BraiNode Chat Transcript\n\n"
        for msg in st.session_state.chat_history:
            role = "User" if msg["role"] == "user" else "BraiNode AI"
            chat_export += f"**{role}:** {msg['content']}\n\n---\n\n"
        st.download_button(
            label="💬 Download Full Chat Transcript (.md)",
            data=chat_export,
            file_name="BraiNode_Chat_History.md",
            mime="text/markdown"
        )

# --- TAB 3: RAGAS EVALUATION ---
with eval_tab:
    st.header("📊 RAGAS Quantitative Benchmarking")
    st.markdown("Scientifically grade your Hybrid RAG pipeline on **Faithfulness** (anti-hallucination) and **Answer Relevancy**.")
    
    if st.button("🚀 Run RAGAS Evaluation on Current Session"):
        if not api_key:
            st.warning("Please enter your Gemini API key in the sidebar first!")
        elif len(st.session_state.eval_data["question"]) == 0:
            st.warning("Please ask at least one question in the Chat tab first!")
        else:
            with st.spinner("Loading RAGAS & grading responses with Gemini... This takes ~30-60 seconds."):
                try:
                    # Lazy-load RAGAS only when clicked so app startup is instant!
                    from datasets import Dataset
                    from ragas import evaluate
                    from ragas.metrics import faithfulness, answer_relevancy
                    from ragas.llms import llm_factory
                    import google.generativeai as genai

                    genai.configure(api_key=api_key)
                    client = genai.GenerativeModel("gemini-1.5-flash")
                    ragas_gemini_llm = llm_factory("gemini-1.5-flash", provider="google", client=client)
                    
                    dataset = Dataset.from_dict(st.session_state.eval_data)
                    score = evaluate(
                        dataset, 
                        metrics=[faithfulness, answer_relevancy],
                        llm=ragas_gemini_llm
                    )
                    
                    st.success("Evaluation Complete!")
                    df_score = score.to_pandas()
                    st.dataframe(df_score, use_container_width=True)
                    
                    csv_data = df_score.to_csv(index=False).encode("utf-8")
                    st.download_button(
                        label="📊 Download Evaluation Metrics (.csv)",
                        data=csv_data,
                        file_name="BraiNode_RAGAS_Metrics.csv",
                        mime="text/csv"
                    )
                except Exception as e:
                    st.error(f"Evaluation Error: {e}")